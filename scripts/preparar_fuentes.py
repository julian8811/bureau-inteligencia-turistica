# -*- coding: utf-8 -*-
"""Descarga fuentes abiertas y escribe solo agregados en data/open.

No guarda microdatos de la EVI ni los XLSX crudos de Aerocivil.
"""
from __future__ import annotations

import csv
import io
import json
import math
import tempfile
import traceback
import zipfile
import unicodedata
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "open"
OUT.mkdir(parents=True, exist_ok=True)
UA = "TablerosBureau/1.0 (investigacion@colmayor.edu.co)"
HOY = datetime.now(timezone.utc).strftime("%Y-%m-%d")

MERCADOS = ["Colombia", "Estados Unidos", "México", "Brasil", "Europa"]
DESTINOS = ["Medellín", "Bogotá", "Cartagena", "Ciudad de México", "São Paulo", "Lima"]
ANIOS = ["2022", "2023", "2024", "2025", "2026"]

EUROPA = {
    "ALBANIA", "ALEMANIA", "ANDORRA", "AUSTRIA", "BELGICA", "BOSNIA", "BOSNIA Y HERZEGOVINA",
    "BULGARIA", "CHIPRE", "CROACIA", "DINAMARCA", "ESLOVAQUIA", "ESLOVENIA", "ESPAÑA", "ESPANA",
    "ESTONIA", "FINLANDIA", "FRANCIA", "GRECIA", "HUNGRIA", "IRLANDA", "ISLANDIA", "ITALIA",
    "LETONIA", "LIECHTENSTEIN", "LITUANIA", "LUXEMBURGO", "MACEDONIA", "MALTA", "MOLDAVIA",
    "MONACO", "MONTENEGRO", "NORUEGA", "PAISES BAJOS", "HOLANDA", "POLONIA", "PORTUGAL",
    "REINO UNIDO", "INGLATERRA", "ESCOCIA", "GALES", "IRLANDA DEL NORTE", "REPUBLICA CHECA",
    "CHECA", "RUMANIA", "RUSIA", "SAN MARINO", "SERBIA", "SUECIA", "SUIZA", "UCRANIA",
    "VATICANO",
}


def fold(text):
    if text is None:
        return ""
    raw = str(text).strip().upper()
    raw = "".join(ch for ch in unicodedata.normalize("NFD", raw) if unicodedata.category(ch) != "Mn")
    return " ".join(raw.split())


def mercado_de(pais):
    nombre = fold(pais)
    if nombre in {"COLOMBIA"}:
        return "Colombia"
    if nombre in {"ESTADOS UNIDOS", "USA", "U.S.A.", "ESTADOS UNIDOS DE AMERICA"}:
        return "Estados Unidos"
    if nombre in {"MEXICO"}:
        return "México"
    if nombre in {"BRASIL", "BRAZIL"}:
        return "Brasil"
    if nombre in EUROPA or nombre.startswith("REINO UNIDO"):
        return "Europa"
    return None


def get(url, timeout=120, headers=None):
    head = {"User-Agent": UA, "Accept": "*/*"}
    if headers:
        head.update(headers)
    req = urllib.request.Request(url, headers=head)
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read(), resp.headers.get("Content-Type", "")


def soda(params):
    url = "https://www.datos.gov.co/resource/gb6w-ynu4.json?" + urllib.parse.urlencode(params)
    raw, _ = get(url, timeout=180, headers={"Accept": "application/json"})
    return json.loads(raw.decode("utf-8"))


def write(name, payload):
    """Reemplaza el JSON solo cuando el texto ya está completo. Un fallo previo deja el archivo anterior."""
    path = OUT / name
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)
    return str(path.relative_to(ROOT)).replace("\\", "/")


def aerocivil():
    llegadas = []
    salidas = []
    for anio in ANIOS:
        llegadas.extend(soda({
            "$select": "a_o,n_mero_de_mes,pais_origen,sum(pasajeros) as pasajeros",
            "$where": f"(destino='MDE' OR destino='EOH') AND a_o='{anio}'",
            "$group": "a_o,n_mero_de_mes,pais_origen",
            "$limit": "50000",
        }))
        salidas.extend(soda({
            "$select": "a_o,n_mero_de_mes,pais_destino,sum(pasajeros) as pasajeros",
            "$where": f"(origen='MDE' OR origen='EOH') AND destino!='MDE' AND destino!='EOH' AND a_o='{anio}'",
            "$group": "a_o,n_mero_de_mes,pais_destino",
            "$limit": "50000",
        }))
    cubo = defaultdict(lambda: {"llegadas": 0, "salidas": 0})
    sin_mapa = defaultdict(int)
    for row in llegadas:
        mercado = mercado_de(row.get("pais_origen"))
        pax = int(float(row.get("pasajeros") or 0))
        if not mercado:
            sin_mapa[row.get("pais_origen") or ""] += pax
            continue
        key = (int(row["a_o"]), int(row["n_mero_de_mes"]), mercado)
        cubo[key]["llegadas"] += pax
    for row in salidas:
        mercado = mercado_de(row.get("pais_destino"))
        pax = int(float(row.get("pasajeros") or 0))
        if not mercado:
            continue
        key = (int(row["a_o"]), int(row["n_mero_de_mes"]), mercado)
        cubo[key]["salidas"] += pax
    serie = []
    for (anio, mes, mercado), vals in sorted(cubo.items()):
        serie.append({
            "anio": anio,
            "mes": mes,
            "mercado": mercado,
            "llegadas": vals["llegadas"],
            "salidas": vals["salidas"],
            "pasajeros": vals["llegadas"] + vals["salidas"],
        })
    return {
        "id": "conectividad",
        "titulo": "Pasajeros comerciales con un extremo en Medellín",
        "fuente": "Aerocivil, Transporte Aéreo Comercial – Tráfico Origen-Destino, datos.gov.co gb6w-ynu4",
        "url": "https://www.datos.gov.co/Transporte/Transporte-A-reo-Comercial-Tr-fico-Origen-Destino-/gb6w-ynu4",
        "licencia": "Datos abiertos de la Aeronáutica Civil, licencia abierta de datos.gov.co",
        "descarga": HOY,
        "universo": "Pasajeros en tráfico comercial con origen o destino José María Córdova (MDE) u Olaya Herrera (EOH). El otro extremo se agrupa en los mercados priorizados.",
        "grano": "mensual",
        "unidad": "pasajeros",
        "nota": "La base publica pasajeros transportados. No trae asientos ofrecidos ni frecuencias.",
        "mercados": MERCADOS,
        "paises_fuera_de_lista": [
            {"pais": k, "llegadas": v} for k, v in sorted(sin_mapa.items(), key=lambda kv: -kv[1])[:12]
        ],
        "serie": serie,
    }


def paises_mensuales():
    """Llegadas y salidas de cada país, por mes. La serie por mercado priorizado no se toca."""
    cubo = defaultdict(lambda: {"llegadas": 0, "salidas": 0})
    for anio in ANIOS:
        llegadas = soda({
            "$select": "a_o,n_mero_de_mes,pais_origen,sum(pasajeros) as pasajeros",
            "$where": f"(destino='MDE' OR destino='EOH') AND a_o='{anio}'",
            "$group": "a_o,n_mero_de_mes,pais_origen",
            "$limit": "50000",
        })
        salidas = soda({
            "$select": "a_o,n_mero_de_mes,pais_destino,sum(pasajeros) as pasajeros",
            "$where": f"(origen='MDE' OR origen='EOH') AND destino!='MDE' AND destino!='EOH' AND a_o='{anio}'",
            "$group": "a_o,n_mero_de_mes,pais_destino",
            "$limit": "50000",
        })
        for row in llegadas:
            pais = fold(row.get("pais_origen"))
            if not pais:
                continue
            cubo[(int(row["a_o"]), int(row["n_mero_de_mes"]), pais)]["llegadas"] += int(float(row.get("pasajeros") or 0))
        for row in salidas:
            pais = fold(row.get("pais_destino"))
            if not pais:
                continue
            cubo[(int(row["a_o"]), int(row["n_mero_de_mes"]), pais)]["salidas"] += int(float(row.get("pasajeros") or 0))
    serie = []
    for (anio, mes, pais), vals in sorted(cubo.items()):
        if not vals["llegadas"] and not vals["salidas"]:
            continue
        mercado = mercado_de(pais)
        serie.append({
            "anio": anio,
            "mes": mes,
            "pais": pais,
            "mercado": mercado,
            "priorizado": mercado is not None,
            "llegadas": vals["llegadas"],
            "salidas": vals["salidas"],
            "pasajeros": vals["llegadas"] + vals["salidas"],
        })
    return serie


def cortes_aerocivil():
    """Aeropuerto, aerolínea y carga. No repite la serie de pasajeros por mercado."""
    aeropuertos = []
    aerolineas = []
    carga = []
    for anio in ANIOS:
        por_apto = soda({
            "$select": "a_o,n_mero_de_mes,destino,sum(pasajeros) as pasajeros",
            "$where": f"(destino='MDE' OR destino='EOH') AND a_o='{anio}'",
            "$group": "a_o,n_mero_de_mes,destino",
            "$limit": "5000",
        })
        for row in por_apto:
            aeropuertos.append({
                "anio": int(row["a_o"]),
                "mes": int(row["n_mero_de_mes"]),
                "aeropuerto": row.get("destino"),
                "llegadas": int(float(row.get("pasajeros") or 0)),
            })
        por_empresa = soda({
            "$select": "a_o,sigla_empresa,nombre,sum(pasajeros) as pasajeros",
            "$where": f"(destino='MDE' OR destino='EOH') AND a_o='{anio}'",
            "$group": "a_o,sigla_empresa,nombre",
            "$limit": "500",
        })
        ordenadas = sorted(por_empresa, key=lambda row: -float(row.get("pasajeros") or 0))[:8]
        for row in ordenadas:
            aerolineas.append({
                "anio": int(row["a_o"]),
                "sigla": row.get("sigla_empresa") or "",
                "nombre": (row.get("nombre") or "").strip(),
                "llegadas": int(float(row.get("pasajeros") or 0)),
            })
        por_carga = soda({
            "$select": "a_o,n_mero_de_mes,sum(carga_correo_kg) as kg",
            "$where": f"(destino='MDE' OR destino='EOH' OR origen='MDE' OR origen='EOH') AND a_o='{anio}'",
            "$group": "a_o,n_mero_de_mes",
            "$limit": "50",
        })
        for row in por_carga:
            carga.append({
                "anio": int(row["a_o"]),
                "mes": int(row["n_mero_de_mes"]),
                "kg": round(float(row.get("kg") or 0), 1),
            })
    return {"aeropuertos": aeropuertos, "aerolineas": aerolineas, "carga": carga}


def wikipedia():
    articulos = {
        "Medellín": ["es.wikipedia", "Medellín", "en.wikipedia", "Medellín"],
        "Bogotá": ["es.wikipedia", "Bogotá", "en.wikipedia", "Bogotá"],
        "Cartagena": ["es.wikipedia", "Cartagena_de_Indias", "en.wikipedia", "Cartagena,_Colombia"],
        "Ciudad de México": ["es.wikipedia", "Ciudad_de_México", "en.wikipedia", "Mexico_City"],
        "São Paulo": ["es.wikipedia", "São_Paulo", "en.wikipedia", "São_Paulo"],
        "Lima": ["es.wikipedia", "Lima", "en.wikipedia", "Lima"],
        "Ciudad de Panamá": ["es.wikipedia", "Ciudad_de_Panamá", "en.wikipedia", "Panama_City"],
        "Santo Domingo": ["es.wikipedia", "Santo_Domingo", "en.wikipedia", "Santo_Domingo"],
        "Quito": ["es.wikipedia", "Quito", "en.wikipedia", "Quito"],
        "Buenos Aires": ["es.wikipedia", "Buenos_Aires", "en.wikipedia", "Buenos_Aires"],
        "Santiago": ["es.wikipedia", "Santiago_de_Chile", "en.wikipedia", "Santiago,_Chile"],
        "Madrid": ["es.wikipedia", "Madrid", "en.wikipedia", "Madrid"],
    }
    serie = []
    usados = []
    for destino, partes in articulos.items():
        for i in range(0, 4, 2):
            proyecto, titulo = partes[i], partes[i + 1]
            path = urllib.parse.quote(titulo, safe="")
            url = (
                "https://wikimedia.org/api/rest_v1/metrics/pageviews/per-article/"
                f"{proyecto}/all-access/user/{path}/monthly/20220101/20261001"
            )
            try:
                raw, _ = get(url, headers={"Accept": "application/json", "Api-User-Agent": UA})
                items = json.loads(raw.decode("utf-8")).get("items", [])
            except Exception:
                continue
            usados.append(f"{proyecto}:{titulo}")
            for item in items:
                stamp = str(item["timestamp"])[:6]
                serie.append({
                    "anio": int(stamp[:4]),
                    "mes": int(stamp[4:6]),
                    "destino": destino,
                    "vistas": int(item["views"]),
                })
    return {
        "id": "wikipedia",
        "fuente": "Wikimedia Pageviews API, proyectos es y en, agente user",
        "url": "https://doc.wikimedia.org/generated-data-platform/aqs/analytics-api/",
        "licencia": "CC0 1.0",
        "descarga": HOY,
        "universo": "Vistas mensuales de los artículos de ciudad. Proxy de interés digital, no de búsquedas en Google.",
        "grano": "mensual, el tablero lo suma al trimestre",
        "articulos": usados,
        "destinos": list(articulos.keys()),
        "serie": serie,
    }


def gdelt():
    import time
    consultas = {
        "Medellín": "Medellin",
        "Bogotá": "Bogota",
        "Cartagena": '"Cartagena de Indias"',
        "Ciudad de México": '"Mexico City"',
        "São Paulo": '"Sao Paulo"',
        "Lima": '"Lima Peru"',
    }
    serie = []
    for destino, query in consultas.items():
        for anio in (2022, 2023, 2024, 2025, 2026):
            fin = "20261001000000" if anio == 2026 else f"{anio + 1}0101000000"
            params = {
                "query": query,
                "mode": "timelinevolraw",
                "format": "json",
                "startdatetime": f"{anio}0101000000",
                "enddatetime": fin,
            }
            url = "https://api.gdeltproject.org/api/v2/doc/doc?" + urllib.parse.urlencode(params)
            try:
                raw, _ = get(url, timeout=25, headers={"Accept": "application/json"})
            except urllib.error.HTTPError as exc:
                if exc.code == 429:
                    raise RuntimeError("GDELT respondió 429 Too Many Requests. No se inventan menciones.") from exc
                raise
            if not raw or not raw[:1] in (b"{", b"["):
                raise RuntimeError("GDELT no entregó JSON para " + destino)
            data = json.loads(raw.decode("utf-8"))
            puntos = []
            for item in data.get("timeline") or []:
                if isinstance(item, dict) and isinstance(item.get("data"), list):
                    puntos.extend(item["data"])
                elif isinstance(item, dict):
                    puntos.append(item)
            for point in puntos:
                fecha = str(point.get("date") or "")
                digitos = "".join(ch for ch in fecha if ch.isdigit())
                valor = point.get("value")
                if len(digitos) < 6 or valor is None:
                    continue
                serie.append({
                    "anio": int(digitos[0:4]),
                    "mes": int(digitos[4:6]),
                    "destino": destino,
                    "menciones": float(valor),
                })
            time.sleep(3)
    if not serie:
        raise RuntimeError("GDELT respondió sin puntos de línea de tiempo")
    return {
        "id": "gdelt",
        "fuente": "GDELT DOC 2.0, timelinevolraw",
        "url": "https://www.gdeltproject.org/",
        "licencia": "Uso abierto con cita a The GDELT Project, https://www.gdeltproject.org/",
        "descarga": HOY,
        "universo": "Volumen de menciones del nombre de la ciudad en medios monitoreados por GDELT. Proxy de share of voice de medios, no de búsqueda.",
        "grano": "el API entrega la línea de tiempo; el tablero suma el año",
        "serie": serie,
    }

MESES_ES = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
    "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12,
}

AEROPUERTOS_MERCADO = {
    "ESTADOS UNIDOS": "MIA",
    "MEXICO": "MEX",
    "BRASIL": "GRU",
    "ESPANA": "MAD",
    "REINO UNIDO": "LHR",
    "FRANCIA": "CDG",
    "ALEMANIA": "FRA",
    "ITALIA": "FCO",
    "PAISES BAJOS": "AMS",
    "HOLANDA": "AMS",
    "PORTUGAL": "LIS",
    "ARGENTINA": "EZE",
    "CHILE": "SCL",
    "PERU": "LIM",
    "ECUADOR": "UIO",
    "PANAMA": "PTY",
    "COSTA RICA": "SJO",
    "CANADA": "YYZ",
    "VENEZUELA": "CCS",
    "REPUBLICA DOMINICANA": "PUJ",
}


def numero(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    text = str(value).strip().replace(" ", "")
    if text.count(",") == 1 and "." not in text:
        text = text.replace(",", ".")
    if text in {"", ".", "NA", "N/A", "-"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def anio_publicado(value):
    """Acepta 2024 y también 2024p, la marca preliminar del anexo."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)
    digitos = "".join(ch for ch in str(value) if ch.isdigit())
    if len(digitos) >= 4:
        return int(digitos[:4])
    return None


def ema():
    import openpyxl
    candidatos = [
        "https://www.dane.gov.co/files/operaciones/EMA/anex-EMA-ago2026.xlsx",
        "https://www.dane.gov.co/files/operaciones/EMA/anex-EMA-dic2025.xlsx",
    ]
    blob, usado = None, None
    for url in candidatos:
        try:
            blob, _ = get(url, timeout=120)
            usado = url
            break
        except Exception:
            continue
    if blob is None:
        raise RuntimeError("No se descargó el anexo EMA")
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    ws = book["6.1 Ind.Var Tarifas.acomoda"]
    rows = list(ws.iter_rows(values_only=True))
    ocupacion = serie_ocupacion(book)
    book.close()
    header_i = next(i for i, row in enumerate(rows) if row and fold(row[0]) == "ANO")
    regiones = []
    for col, cell in enumerate(rows[header_i]):
        if col < 2 or cell is None or not str(cell).strip():
            continue
        if fold(cell).startswith("TOTAL"):
            continue
        regiones.append((col, str(cell).strip()))
    serie = []
    anio = None
    for row in rows[header_i + 1:]:
        if not row:
            continue
        if row[0] not in (None, ""):
            anio_num = anio_publicado(row[0])
            if anio_num:
                anio = anio_num
        mes = MESES_ES.get(fold(row[1] if len(row) > 1 else ""))
        if not anio or not mes or anio < 2022:
            continue
        niveles = {}
        for col, nombre in regiones:
            if col >= len(row):
                continue
            valor = numero(row[col])
            if valor is not None:
                niveles[nombre] = valor
        anti = next((v for k, v in niveles.items() if fold(k).startswith("ANTIOQUIA")), None)
        if anti is None or len(niveles) < 3:
            continue
        promedio = sum(niveles.values()) / len(niveles)
        serie.append({
            "anio": anio,
            "mes": mes,
            "indice_antioquia": round(anti, 2),
            "indice_promedio_regiones": round(promedio, 2),
            "indice_100": round(anti / promedio * 100, 2),
        })
    return {
        "id": "ema",
        "fuente": "DANE, Encuesta Mensual de Alojamiento, anexo hoja 6.1, índice de tarifa sencilla",
        "url": usado,
        "licencia": "Estadística oficial DANE",
        "descarga": HOY,
        "universo": "Regiones EMA. Antioquia no es la tarifa de la ciudad de Medellín.",
        "grano": "mensual",
        "base": "Índice DANE base 2019=100, acomodación sencilla. indice_100 compara Antioquia con el promedio de las regiones de la fila, sin el total nacional.",
        "ocupacion": ocupacion,
        "serie": serie,
    }


def serie_ocupacion(book):
    """Porcentaje mensual de ocupación por región. Las hojas de ingresos y personal del anexo son variaciones, no niveles."""
    nombre = next((n for n in book.sheetnames if n.startswith("4.2")), None)
    if not nombre:
        return {"estado": "El anexo no trae la hoja 4.2 de ocupación mensual.", "serie": []}
    rows = list(book[nombre].iter_rows(values_only=True))
    header_i = next(i for i, row in enumerate(rows) if row and fold(row[0]) == "ANO")
    header = rows[header_i]
    nombres = []
    for i, cell in enumerate(header):
        if i < 2 or cell is None or not str(cell).strip():
            continue
        nombres.append((i, str(cell).strip()))
    if not any("ANTIOQUIA" in fold(nombre) for _, nombre in nombres):
        return {"estado": "La hoja 4.2 no tiene columna de Antioquia.", "serie": []}
    serie = []
    anio = None
    for row in rows[header_i + 1:]:
        if not row:
            continue
        if row[0] not in (None, ""):
            hallado = anio_publicado(row[0])
            if hallado:
                anio = hallado
        mes = MESES_ES.get(fold(row[1] if len(row) > 1 else ""))
        if not anio or not mes or anio < 2022:
            continue
        regiones = {}
        for col, nombre in nombres:
            if col >= len(row):
                continue
            valor = numero(row[col])
            if valor is not None:
                regiones[nombre] = round(valor, 2)
        if not any(fold(k).startswith("ANTIOQUIA") for k in regiones):
            continue
        punto = {"anio": anio, "mes": mes, "regiones": regiones}
        for clave, fragmento in (("total_nacional", "TOTAL"), ("bogota", "BOGOTA"), ("cartagena", "CARTAGENA"), ("antioquia", "ANTIOQUIA")):
            hallado = next((v for k, v in regiones.items() if fragmento in fold(k)), None)
            if hallado is not None:
                punto[clave] = hallado
        serie.append(punto)
    return {
        "hoja": nombre,
        "unidad": "porcentaje de ocupación",
        "universo": "Regiones EMA. Antioquia no es la ocupación de la ciudad de Medellín.",
        "serie": serie,
    }


def medata():
    paginas = [
        "https://medata.app.medellin.gov.co/node/16279",
        "https://www.medata.gov.co/node/16279",
    ]
    html, pagina = None, None
    ultimo = None
    for url in paginas:
        try:
            html, _ = get(url, timeout=40)
            pagina = url
            break
        except Exception as exc:
            ultimo = exc
    if html is None:
        raise RuntimeError("Medata no respondió: " + str(ultimo))
    texto = html.decode("utf-8", "replace")
    url = None
    for token in texto.replace("'", '"').split('"'):
        bajo = token.lower()
        if bajo.endswith(".csv"):
            url = token
            break
    if url and url.startswith("/"):
        url = "https://www.medata.gov.co" + url
    if not url:
        raise RuntimeError("Medata no expuso un CSV en la ficha de tarifa por zona")
    blob, _ = get(url, timeout=60)
    lineas = blob.decode("utf-8-sig", "replace").splitlines()
    cab = [c.strip() for c in lineas[0].split(",")]
    serie = []
    for linea in lineas[1:]:
        if not linea.strip():
            continue
        partes = [p.strip() for p in linea.split(",")]
        serie.append({cab[i] if i < len(cab) else str(i): partes[i] for i in range(min(len(cab), len(partes)))})
    return {
        "id": "medata_tarifas",
        "fuente": "Medata, tarifa promedio mensual por zona de Medellín",
        "url": pagina,
        "archivo": url,
        "licencia": "CC BY-SA 4.0",
        "descarga": HOY,
        "corte_publicacion": "2023-12-14",
        "universo": "Tarifa promedio mensual por zona. El metadato de la ficha marca la última modificación el 14 de diciembre de 2023.",
        "grano": "mensual",
        "serie": serie[:360],
    }


def ola_evi(anio, mes):
    clave = anio * 100 + mes
    if 202308 <= clave <= 202407:
        return "Agosto 2023 – julio 2024"
    if 202501 <= clave <= 202506:
        return "Enero – junio 2025"
    return None


def es_jmc(aeropuerto):
    nombre = fold(aeropuerto)
    return nombre == "MEDELLIN" or "RIONEGRO" in nombre or "CORDOVA" in nombre or "JOSE MAR" in nombre


def filas_zip(url):
    raw, _ = get(url, timeout=300)
    fd, nombre = tempfile.mkstemp(suffix=".zip")
    path = Path(nombre)
    with open(fd, "wb") as handle:
        handle.write(raw)
    try:
        with zipfile.ZipFile(path) as zf:
            csv_name = next(n for n in zf.namelist() if n.lower().endswith(".csv"))
            with zf.open(csv_name) as raw_csv:
                text = io.TextIOWrapper(raw_csv, encoding="utf-8-sig", newline="")
                for row in csv.DictReader(text):
                    yield row
    finally:
        path.unlink(missing_ok=True)


def top_pesos(cubo, limite):
    total = sum(cubo.values()) or 1
    orden = sorted(cubo.items(), key=lambda kv: -kv[1])[:limite]
    return [{"nombre": k, "peso": round(v, 1), "participacion": round(v / total * 100, 1)} for k, v in orden]


def con_peso(cubo, total, minimo=0.005):
    """Países cuya participación en el peso de la ola alcanza el umbral. El resto no se publica como promedio."""
    base = total or 1
    orden = sorted(cubo.items(), key=lambda kv: -kv[1])
    return [
        {"nombre": k, "peso": round(v, 1), "participacion": round(v / base * 100, 1)}
        for k, v in orden
        if v / base >= minimo
    ]


def agregar_evi():
    url = "https://microdatos.dane.gov.co/index.php/catalog/869/download/24177"
    olas = {}
    meses = {}
    aeropuertos = defaultdict(int)

    def cubo():
        return {
            "n": 0,
            "peso": 0.0,
            "gasto": 0.0,
            "gasto_w": 0.0,
            "diario": 0.0,
            "diario_w": 0.0,
            "noches": 0.0,
            "noches_w": 0.0,
            "grupo": defaultdict(float),
            "motivo": defaultdict(float),
            "residencia": defaultdict(float),
            "mercado": defaultdict(float),
            "gasto_mercado": defaultdict(lambda: {"x": 0.0, "w": 0.0}),
            "gasto_pais": defaultdict(lambda: {"x": 0.0, "w": 0.0}),
            "distancia": defaultdict(float),
        }

    for row in filas_zip(url):
        aero = row.get("Aeropuerto_puerto") or ""
        aeropuertos[aero] += 1
        if not es_jmc(aero):
            continue
        if fold(row.get("P3807")) != "NO":
            continue
        mes = MESES_ES.get(fold(row.get("Mes")))
        anio = numero(row.get("Vigencia"))
        if not mes or not anio:
            continue
        nombre_ola = ola_evi(int(anio), mes)
        if not nombre_ola:
            continue
        peso = numero(row.get("fexp")) or 0
        if peso <= 0:
            continue
        slot = olas.setdefault(nombre_ola, cubo())
        slot["n"] += 1
        slot["peso"] += peso
        clave_mes = (int(anio), mes)
        mes_slot = meses.setdefault(clave_mes, {"n": 0, "gasto": 0.0, "gasto_w": 0.0, "noches": 0.0, "noches_w": 0.0})
        mes_slot["n"] += 1
        grupo = numero(row.get("P3835"))
        gasto = numero(row.get("P3810"))
        noches = numero(row.get("Noches_per_prin"))
        if grupo and 1 <= grupo <= 30:
            if grupo == 1:
                etiqueta = "Solo"
            elif grupo == 2:
                etiqueta = "Pareja"
            elif grupo <= 4:
                etiqueta = "Grupo de 3 a 4"
            else:
                etiqueta = "Grupo de 5 o más"
            slot["grupo"][etiqueta] += peso
        if noches and 1 <= noches <= 90:
            slot["noches"] += noches * peso
            slot["noches_w"] += peso
            mes_slot["noches"] += noches * peso
            mes_slot["noches_w"] += peso
        if gasto and grupo and 1 <= grupo <= 30 and 1 <= gasto / grupo <= 20000:
            por_persona = gasto / grupo
            slot["gasto"] += por_persona * peso
            slot["gasto_w"] += peso
            mes_slot["gasto"] += por_persona * peso
            mes_slot["gasto_w"] += peso
            if noches and 1 <= noches <= 90:
                slot["diario"] += (por_persona / noches) * peso
                slot["diario_w"] += peso
        motivo = (row.get("P3821") or "").strip()
        if motivo:
            slot["motivo"][motivo] += peso
        pais = (row.get("P3808") or "").strip()
        if pais:
            slot["residencia"][pais] += peso
            mercado = mercado_de(pais)
            if gasto and grupo and 1 <= grupo <= 30 and 1 <= gasto / grupo <= 20000:
                slot["gasto_pais"][pais]["x"] += (gasto / grupo) * peso
                slot["gasto_pais"][pais]["w"] += peso
                if mercado:
                    slot["gasto_mercado"][mercado]["x"] += (gasto / grupo) * peso
                    slot["gasto_mercado"][mercado]["w"] += peso
            if mercado:
                slot["mercado"][mercado] += peso
            slot["distancia"][fold(pais)] += peso
    salida = []
    for nombre, slot in olas.items():
        salida.append({
            "ola": nombre,
            "n": slot["n"],
            "peso": round(slot["peso"], 1),
            "gasto_estadia_usd": round(slot["gasto"] / slot["gasto_w"], 2) if slot["gasto_w"] else None,
            "gasto_diario_usd": round(slot["diario"] / slot["diario_w"], 2) if slot["diario_w"] else None,
            "noches": round(slot["noches"] / slot["noches_w"], 2) if slot["noches_w"] else None,
            "n_gasto": round(slot["gasto_w"], 1),
            "n_noches": round(slot["noches_w"], 1),
            "grupo": top_pesos(slot["grupo"], 6),
            "motivos": top_pesos(slot["motivo"], 8),
            "residencia": con_peso(slot["residencia"], slot["peso"]),
            "umbral_pais": "0,5 % del peso de la ola",
            "mercados": top_pesos(slot["mercado"], 6),
            "gasto_por_mercado": [
                {"mercado": k, "gasto_estadia_usd": round(v["x"] / v["w"], 2), "peso": round(v["w"], 1)}
                for k, v in slot["gasto_mercado"].items() if v["w"]
            ],
            "gasto_por_pais": [
                {"pais": k, "gasto_estadia_usd": round(v["x"] / v["w"], 2), "peso": round(v["w"], 1)}
                for k, v in slot["gasto_pais"].items()
                if v["w"] and slot["peso"] and v["w"] / slot["peso"] >= 0.005
            ],
            "paises_peso": {k: round(v, 1) for k, v in slot["distancia"].items()},
        })
    return {
        "id": "evi",
        "fuente": "DANE, Encuesta de Visitantes Internacionales, microdatos catálogo 869, base aéreo-marítimo",
        "url": "https://microdatos.dane.gov.co/index.php/catalog/869",
        "archivo": url,
        "licencia": "Microdatos DANE. El tablero guarda promedios y distribuciones ponderadas con fexp. Las filas no se publican.",
        "descarga": HOY,
        "universo": "No residente encuestado en el aeropuerto José María Córdova (Rionegro). Las noches son las del destino principal del viaje (Noches_per_prin). El gasto está en dólares por persona del grupo.",
        "grano": "ola publicada",
        "aeropuertos_en_la_base": [{"aeropuerto": k, "filas": v} for k, v in sorted(aeropuertos.items(), key=lambda kv: -kv[1])],
        "meses": [
            {
                "anio": anio,
                "mes": mes,
                "n": slot["n"],
                "gasto_estadia_usd": round(slot["gasto"] / slot["gasto_w"], 2) if slot["gasto_w"] else None,
                "noches": round(slot["noches"] / slot["noches_w"], 2) if slot["noches_w"] else None,
            }
            for (anio, mes), slot in sorted(meses.items())
        ],
        "olas": salida,
    }


def haversine(lat1, lon1, lat2, lon2):
    r = 6371.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlmb = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlmb / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def coordenadas_aeropuertos():
    raw, _ = get("https://davidmegginson.github.io/ourairports-data/airports.csv", timeout=120)
    wanted = set(AEROPUERTOS_MERCADO.values()) | {"MDE"}
    puntos = {}
    text = io.StringIO(raw.decode("utf-8"))
    for row in csv.DictReader(text):
        iata = (row.get("iata_code") or "").strip()
        if iata in wanted and iata not in puntos:
            puntos[iata] = (float(row["latitude_deg"]), float(row["longitude_deg"]), row.get("name"))
    return puntos


def factor_kg(row):
    """El libro plano trae un identificador enorme y, al final, el factor en kg CO2e."""
    candidatos = []
    for celda in row:
        valor = numero(celda)
        if valor is not None and 0 < valor < 10000:
            candidatos.append(valor)
    return candidatos[-1] if candidatos else None


def factores_defra():
    import openpyxl
    pagina = "https://www.gov.uk/government/publications/greenhouse-gas-reporting-conversion-factors-2026"
    html, _ = get(pagina, timeout=60)
    texto = html.decode("utf-8", "replace")
    hrefs = []
    cursor = 0
    while True:
        pos = texto.lower().find(".xlsx", cursor)
        if pos < 0:
            break
        inicio = texto.rfind("http", max(0, pos - 300), pos)
        if inicio >= 0:
            hrefs.append(texto[inicio:pos + 5].split('"')[0].split("'")[0])
        cursor = pos + 5
    if not hrefs:
        raise RuntimeError("La página de factores 2026 no mostró un xlsx")
    preferido = next((h for h in hrefs if "flat" in h.lower()), hrefs[0])
    blob, _ = get(preferido, timeout=180)
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    aereo = None
    hotel = None
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            celdas = [str(c) for c in row if c is not None]
            if not celdas:
                continue
            junto = " | ".join(celdas)
            plano = fold(junto)
            factor = factor_kg(row)
            if aereo is None and factor and "LONG-HAUL" in plano and "AVERAGE" in plano and "PASSENGER" in plano and "KG CO2E" in plano:
                aereo = {"factor": factor, "fila": junto[:240], "hoja": sheet.title}
            if hotel is None and factor and "HOTEL" in plano and "COLOMBIA" in plano and "KG CO2E" in plano:
                hotel = {"factor": factor, "fila": junto[:240], "hoja": sheet.title}
            if aereo and hotel:
                break
        if aereo and hotel:
            break
    book.close()
    if not aereo or not hotel:
        raise RuntimeError(f"No se aislaron ambos factores. Aéreo={bool(aereo)} hotel={bool(hotel)} archivo={preferido}")
    return {"url": preferido, "aereo_kg_por_pasajero_km": aereo, "hotel_kg_por_noche": hotel}


def huella_desde(evi):
    puntos = coordenadas_aeropuertos()
    if "MDE" not in puntos:
        raise RuntimeError("OurAirports no trajo el código MDE")
    factores = factores_defra()
    mde = puntos["MDE"]
    distancias = {}
    for pais, iata in AEROPUERTOS_MERCADO.items():
        if iata not in puntos:
            continue
        distancias[pais] = round(haversine(puntos[iata][0], puntos[iata][1], mde[0], mde[1]), 1)
    serie = []
    for ola in evi.get("olas", []):
        pesos = ola.get("paises_peso") or {}
        num = den = 0
        for pais, peso in pesos.items():
            km = distancias.get(fold(pais))
            if km is None:
                continue
            num += km * peso
            den += peso
        km_medio = num / den if den else None
        noches = ola.get("noches")
        factor_aire = factores["aereo_kg_por_pasajero_km"]["factor"]
        factor_hotel = factores["hotel_kg_por_noche"]["factor"]
        vuelo = round(km_medio * factor_aire, 2) if km_medio is not None else None
        aloj = round(noches * factor_hotel, 2) if noches is not None else None
        total = round(vuelo + aloj, 2) if vuelo is not None and aloj is not None else None
        serie.append({
            "ola": ola["ola"],
            "km_ida": round(km_medio, 1) if km_medio is not None else None,
            "kg_vuelo": vuelo,
            "kg_alojamiento": aloj,
            "kg_total": total,
            "cobertura_peso": round(den, 1),
        })
    return {
        "id": "huella",
        "fuente": "Distancias OurAirports hasta MDE, noches de la EVI y factores de conversión del gobierno del Reino Unido 2026",
        "url": factores["url"],
        "licencia": "Los factores se publican para el reporte de GEI. El uso fuera de ese reporte queda bajo responsabilidad de quien los aplica.",
        "descarga": HOY,
        "nota": "Tramo de ida, no el regreso. El factor de hotel se aplica por noche de persona, sin observar si el grupo comparte habitación. Estimación, no medición directa.",
        "factores": factores,
        "aeropuerto_mde": {"lat": mde[0], "lon": mde[1], "nombre": mde[2]},
        "distancias_km": distancias,
        "serie": serie,
    }


def evi_y_huella(errores):
    try:
        evi = agregar_evi()
    except Exception:
        errores.append(traceback.format_exc(limit=3))
        evi = {
            "id": "evi",
            "fuente": "DANE, Encuesta de Visitantes Internacionales, catálogo 869",
            "url": "https://microdatos.dane.gov.co/index.php/catalog/869",
            "descarga": HOY,
            "estado": "No se pudo agregar la EVI en esta corrida.",
            "olas": [],
        }
        return evi, {
            "id": "huella",
            "estado": "Sin agregados EVI no hay huella.",
            "serie": [],
            "descarga": HOY,
        }
    try:
        huella = huella_desde(evi)
    except Exception:
        errores.append(traceback.format_exc(limit=3))
        huella = {
            "id": "huella",
            "fuente": "EVI y factores del gobierno del Reino Unido",
            "descarga": HOY,
            "estado": "La EVI quedó agregada y los factores o las distancias no se pudieron aislar.",
            "serie": [],
        }
    for ola in evi.get("olas", []):
        ola.pop("paises_peso", None)
    return evi, huella


def icca():
    return {
        "id": "icca",
        "fuente": "ICCA GlobeWatch Country and City Rankings 2025, tabla de ciudades",
        "url": "https://wot.waw.pl/wp-content/uploads/2026/05/ICCA-GlobeWatch-Business-Analytics-%E2%80%93-Country-and-City-Rankings-2025.pdf",
        "licencia": "Cifras citadas del ranking difundido en ese PDF. La base ICCA Business Intelligence sigue siendo de miembros.",
        "descarga": HOY,
        "anio": 2025,
        "grano": "anual",
        "nota": "Reuniones de asociación elegibles ICCA. No hay participación no ICCA ni corte trimestral en esta fuente.",
        "ciudades": [
            {"ciudad": "Bogotá", "puesto": 36, "reuniones": 63},
            {"ciudad": "Cartagena", "puesto": 77, "reuniones": 36},
            {"ciudad": "Medellín", "puesto": 103, "reuniones": 26},
        ],
    }


def main():
    errores = []
    piezas = []
    for nombre, fn in (
        ("conectividad.json", aerocivil),
        ("wikipedia.json", wikipedia),
        ("gdelt.json", gdelt),
        ("ema.json", ema),
        ("medata_tarifas.json", medata),
        ("icca.json", icca),
    ):
        try:
            payload = fn()
            piezas.append({"archivo": write(nombre, payload), "id": payload.get("id"), "estado": "ok", "n": len(payload.get("serie") or payload.get("ciudades") or [])})
            print("ok", nombre, piezas[-1]["n"])
        except Exception:
            detalle = traceback.format_exc(limit=2)
            errores.append(f"{nombre}: {detalle}")
            print("fallo", nombre)
            print(detalle)
    evi, huella = evi_y_huella(errores)
    piezas.append({"archivo": write("evi.json", evi), "id": "evi", "estado": evi.get("estado", "ok"), "n": len(evi.get("olas") or [])})
    piezas.append({"archivo": write("huella.json", huella), "id": "huella", "estado": huella.get("estado", "ok"), "n": len(huella.get("serie") or [])})
    print("evi", piezas[-2])
    print("huella", piezas[-1])
    manifiesto = {
        "generado": HOY,
        "piezas": piezas,
        "errores": errores,
        "mercados": MERCADOS,
        "destinos_comparables": DESTINOS,
        "nota": "Agregados para el tablero. Sin filas de encuesta y sin XLSX crudos.",
    }
    write("manifiesto.json", manifiesto)
    print("manifiesto listo")


if __name__ == "__main__":
    main()
